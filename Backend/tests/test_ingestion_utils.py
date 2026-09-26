import io
import unittest
from unittest.mock import patch

from fastapi import HTTPException
from docx import Document

from app.routes.upload_file import ingest_document, normalize_doc_id
from app.utils.preprocess_text import (
    chunk_sections,
    chunk_text,
    extract_sections,
    extract_text,
)


class IngestionUtilityTests(unittest.TestCase):
    def test_document_id_is_sanitized(self):
        self.assertEqual(normalize_doc_id(" Admissions / 2026 "), "Admissions-2026")

    def test_invalid_document_id_is_rejected(self):
        with self.assertRaises(HTTPException):
            normalize_doc_id("////")

    def test_chunks_have_overlap_without_duplicate_tail(self):
        text = " ".join(str(number) for number in range(10))
        chunks = chunk_text(text, chunk_size=5, overlap=2)
        self.assertEqual(chunks, ["0 1 2 3 4", "3 4 5 6 7", "6 7 8 9"])

    def test_txt_must_be_utf8(self):
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            extract_text(b"\xff\xfe", "document.txt")

    def test_docx_extracts_tables_in_order_with_section_titles(self):
        document = Document()
        document.add_heading("Overview", level=1)
        document.add_paragraph("BUP is a public university.")
        document.add_heading("Key Statistics", level=1)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Students"
        table.cell(0, 1).text = "6,948"
        document.add_heading("Mission", level=1)
        document.add_paragraph("Serve the nation.")
        file_buffer = io.BytesIO()
        document.save(file_buffer)
        file_bytes = file_buffer.getvalue()

        extracted = extract_text(file_bytes, "document.docx")
        self.assertLess(extracted.index("Overview"), extracted.index("Students | 6,948"))
        self.assertLess(extracted.index("Students | 6,948"), extracted.index("Mission"))

        sections = extract_sections(file_bytes, "document.docx")
        section_chunks = chunk_sections(sections)
        statistics_chunks = [
            chunk for chunk, title in section_chunks if title == "Key Statistics"
        ]
        self.assertEqual(len(statistics_chunks), 1)
        self.assertIn("Key Statistics: Students | 6,948", statistics_chunks[0])

    def test_successful_ingestion_clears_answer_cache(self):
        with (
            patch("app.routes.upload_file.get_embedding", return_value=["embedding"]),
            patch("app.routes.upload_file.store_embeddings", return_value=1),
            patch("app.routes.upload_file.redis_cache_service.clear_cache") as clear_cache,
        ):
            stored_count = ingest_document(b"BUP information", "info.txt", "bup-info")

        self.assertEqual(stored_count, 1)
        clear_cache.assert_called_once_with()

    def test_failed_ingestion_does_not_clear_answer_cache(self):
        with (
            patch("app.routes.upload_file.get_embedding", return_value=["embedding"]),
            patch(
                "app.routes.upload_file.store_embeddings",
                side_effect=RuntimeError("Pinecone unavailable"),
            ),
            patch("app.routes.upload_file.redis_cache_service.clear_cache") as clear_cache,
        ):
            with self.assertRaisesRegex(RuntimeError, "Pinecone unavailable"):
                ingest_document(b"BUP information", "info.txt", "bup-info")

        clear_cache.assert_not_called()


if __name__ == "__main__":
    unittest.main()
